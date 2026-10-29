#!/usr/bin/env python
"""把荔园畅课待办写成一份 Excel 表，并保证反复运行只维护同一个文件。

维护规则（与需求一致）:
  - 表格路径固定，每次重跑重写同一个文件，不新建、不追加副本。
  - 行序永远按 DDL 升序：
      旧待办从系统消失 -> 整行删除，下方待办整体上移
      出现新待办       -> 按 DDL 排序插入到对应位置
  - 只含自动字段，无人工列。
"""

from __future__ import annotations

import os
import time
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import core_lock

SHEET_NAME = "待办"

HEADERS = ["序号", "任务名称", "课程", "截止时间", "剩余", "附件", "附件文件", "链接"]
WIDTHS = [6, 46, 26, 18, 13, 7, 52, 8]
ATT_NAME_WIDTH = 44  # 附件文件名列按此显示宽度硬折行

TYPE_CN = {
    "homework": "作业",
    "quiz": "测验",
    "exam": "考试",
    "material": "资料",
    "discussion": "讨论",
    "attendance": "考勤",
    "questionnaire": "问卷",
}

# ---- 样式 ----
HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
OVERDUE_FILL = PatternFill("solid", fgColor="FFC7CE")
URGENT_FILL = PatternFill("solid", fgColor="FDE2E2")   # < 24h
SOON_FILL = PatternFill("solid", fgColor="FFF2CC")     # < 72h
STRIPE_FILL = PatternFill("solid", fgColor="F5F8FD")
THIN = Side(style="thin", color="D6DCE5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fmt_left(seconds: float) -> str:
    if seconds < 0:
        sec = -seconds
        days = sec / 86400
        return f"已过期 {days:.0f} 天" if days >= 1 else f"已过期 {int(sec // 3600)} 小时"
    if seconds < 3600:
        return f"{int(seconds // 60)} 分钟"
    if seconds < 86400:
        return f"{seconds / 3600:.1f} 小时"
    return f"{seconds / 86400:.1f} 天"


def _urgency(seconds: float | None):
    if seconds is None:
        return None
    if seconds < 0:
        return OVERDUE_FILL
    if seconds < 86400:
        return URGENT_FILL
    if seconds < 86400 * 3:
        return SOON_FILL
    return None


def _key(item: dict) -> tuple[str, str]:
    return (item.get("course_name", "").strip(), item.get("title", "").strip())


def _dwidth(s: str) -> int:
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in s)


def _hard_wrap(text: str, width: int) -> str:
    """按显示宽度硬折行，优先在「、」后断开，保证 Excel 不会再自行折行。"""
    text = (text or "").replace("\r", "").replace("\n", "、")
    if not text:
        return ""
    lines: list[str] = []
    cur = ""
    for part in text.split("、"):
        candidate = part if not cur else f"{cur}、{part}"
        if _dwidth(candidate) <= width:
            cur = candidate
            continue
        if cur:
            lines.append(cur)
        cur = ""
        # 单段本身超宽时按字符硬切
        piece = ""
        for ch in part:
            if _dwidth(piece + ch) > width and piece:
                lines.append(piece)
                piece = ch
            else:
                piece += ch
        cur = piece
    if cur:
        lines.append(cur)
    return "\n".join(lines)


def _wrap_lines(text: str, col_width: int) -> int:
    """估算文本在给定列宽下折成几行（CJK 字符按 2 个字符宽计）。"""
    if not text:
        return 1
    lines, used = 1, 0
    for ch in str(text):
        w = 2 if unicodedata.east_asian_width(ch) in "WF" else 1
        if used + w > col_width:
            lines += 1
            used = w
        else:
            used += w
    return lines


def read_existing(path: Path) -> dict[tuple[str, str], dict]:
    """读旧表，返回 {(课程, 任务): {附件数, 附件文件}}。文件不存在或损坏时返回空。"""
    out: dict[tuple[str, str], dict] = {}
    path = Path(path)
    if not path.exists():
        return out
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return out
    try:
        if SHEET_NAME not in wb.sheetnames:
            return out
        ws = wb[SHEET_NAME]
        rows = ws.iter_rows(values_only=True)
        next(rows, None)  # 跳过表头
        for row in rows:
            if not row or len(row) < 7:
                continue
            title, course = row[1], row[2]
            if not title:
                continue
            out[(str(course or "").strip(), str(title).strip())] = {
                "att_count": row[5],
                "att_files": row[6],
            }
    finally:
        wb.close()
    return out


def _fill_sheet(wb: Workbook, items: list[dict], now: datetime):
    """在 wb 中重建「待办」sheet（原位置替换），返回 worksheet。"""
    if SHEET_NAME in wb.sheetnames:
        pos = wb.sheetnames.index(SHEET_NAME)
        del wb[SHEET_NAME]
    else:
        pos = 0
    ws = wb.create_sheet(SHEET_NAME, pos)

    for i, name in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[c.column_letter].width = WIDTHS[i - 1]
    ws.row_dimensions[1].height = 26

    for idx, item in enumerate(items, 1):
        r = idx + 1
        deadline = datetime.fromisoformat(item["deadline"]) if item.get("deadline") else None
        seconds = (deadline - now).total_seconds() if deadline else None

        if item.get("attachments_skipped"):
            att_count = item.get("old_att_count") or 0
            names = item.get("old_att_files") or ""
        else:
            att_list = item.get("attachments") or []
            att_count = len(att_list)
            names = "、".join(a.get("name", "") for a in att_list if a.get("name"))

        names = _hard_wrap(names, ATT_NAME_WIDTH)

        values = [
            idx,
            item.get("title", ""),
            item.get("course_name", ""),
            deadline.replace(tzinfo=None) if deadline else "—",
            _fmt_left(seconds) if seconds is not None else "—",
            att_count,
            names,
            item.get("url", ""),
        ]
        for i, v in enumerate(values, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = BODY_FONT
            c.border = BORDER
            if i in (2, 3, 7):
                c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                c.alignment = Alignment(horizontal="center", vertical="center")

        if deadline:
            ws.cell(row=r, column=4).number_format = "yyyy-mm-dd hh:mm"

        # 按折行数显式设置行高，避免长任务名/长附件名压行
        lines = max(
            _wrap_lines(item.get("title", ""), WIDTHS[1] - 2),
            _wrap_lines(item.get("course_name", ""), WIDTHS[2] - 2),
            names.count("\n") + 1 if names else 1,
        )
        ws.row_dimensions[r].height = max(19, lines * 15.5 + 4)

        link_cell = ws.cell(row=r, column=8)
        if item.get("url"):
            link_cell.value = "打开"
            link_cell.hyperlink = item["url"]
            link_cell.font = Font(name="微软雅黑", size=10, color="0563C1", underline="single")

        fill = _urgency(seconds)
        overdue = seconds is not None and seconds < 0
        if fill is None and idx % 2 == 0:
            fill = STRIPE_FILL
        if fill is not None:
            for i in range(1, len(HEADERS) + 1):
                ws.cell(row=r, column=i).fill = fill
        if overdue:
            for i in (2, 3, 4, 5):
                ws.cell(row=r, column=i).font = Font(name="微软雅黑", size=10, color="9C0006", bold=True)

    ws.freeze_panes = "A2"
    if items:
        ws.auto_filter.ref = f"A1:{chr(64 + len(HEADERS))}{len(items) + 1}"
    ws.sheet_view.showGridLines = False
    return ws


def _drop_blank_sheets(wb: Workbook):
    """删掉新建工作簿残留的空白默认页（不动有内容的 sheet）。"""
    for name in list(wb.sheetnames):
        if name == SHEET_NAME or len(wb.sheetnames) <= 1:
            continue
        ws = wb[name]
        if ws.max_row <= 1 and ws.max_column <= 1 and ws["A1"].value is None:
            del wb[name]


def update_sheet(path: Path, items: list[dict], now: datetime, log=print,
                 close_preview: bool = True) -> dict:
    """按 DDL 重写整张表，返回变更摘要。

    items 必须已按 DDL 升序排好；本函数不再排序（排序由调用方保证）。
    close_preview=True 时，若表格被 WorkBuddy 预览内核占用会自动关掉它再写；
    被 Excel / WPS 占用则只报错不强制关闭（避免丢未保存的修改）。
    """
    path = Path(path)
    old = read_existing(path)
    new_keys = {_key(it) for it in items}
    added = [it for it in items if _key(it) not in old]
    removed = [k for k in old if k not in new_keys]

    if path.exists():
        try:
            wb = load_workbook(path)
        except Exception:
            wb = Workbook()
    else:
        wb = Workbook()

    _fill_sheet(wb, items, now)
    _drop_blank_sheets(wb)
    wb.active = 0

    tmp = path.with_suffix(path.suffix + ".tmp")
    try:
        wb.save(tmp)
    except Exception as exc:
        _cleanup(tmp)
        log(f"[!] 写表格失败: {exc}")
        return {"error": str(exc), "added": added, "removed": removed}

    # 写入目标：占用时先尝试释放（预览内核），再重试
    written = False
    reason: str | None = None
    for attempt in range(4):
        try:
            os.replace(tmp, path)
            written = True
            break
        except PermissionError:
            if attempt == 0:
                reason = core_lock.release(path, log, close_preview=close_preview)
            time.sleep(0.35)
        except Exception as exc:
            _cleanup(tmp)
            log(f"[!] 写表格失败: {exc}")
            return {"error": str(exc), "added": added, "removed": removed}

    if not written:
        _cleanup(tmp)
        msg = reason or f"无法写入 {path.name}，请稍后重试"
        log(f"[!] {msg}")
        return {"error": msg, "added": added, "removed": removed}

    return {"ok": True, "added": added, "removed": removed, "total": len(items)}


def _cleanup(tmp: Path):
    try:
        tmp.unlink(missing_ok=True)
    except Exception:
        pass
